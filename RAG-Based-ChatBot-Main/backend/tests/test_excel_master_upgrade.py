"""
DocuChat AI - Master Excel Upgrade Comprehensive Test Suite
Validates:
1. Benchmark Workbook Generation: Sales Performance Analysis Dashboard — 2026
2. 4 Primary Sheets: Dashboard, Sales Data, Analysis, Lists
3. 30 Transactions, Unique IDs, Indian Rupee (₹) formatting, Date formatting
4. Working Excel Formulas: =F2*G2, =H2-I2, =IFERROR(J2/H2,0), SUM, AVERAGE, COUNT
5. 5 Native Openpyxl Charts (Region Bar, Product Column, Region Pie, Monthly Line, Salesperson Bar)
6. Data Validation rules (Dropdowns from Lists) and Conditional Formatting rules
7. Independent Python calculation reconciliation and formula error auditing
8. Advanced Excel Analyzer: Top-N ranking, monthly summaries, audit, natural language QA
9. AI Agent Routing: Generation vs Spreadsheet Analysis vs Formula Explainer vs PDF RAG vs Web Search
10. Security: CWE-1236 formula injection protection preserving safe formulas
"""

import os
import json
import pytest
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

from backend.config import settings
from backend.tools.excel_generator import (
    sanitize_cell_value,
    generate_workbook_from_request,
    active_excel_artifact,
)
from backend.tools.benchmark_dashboard_2026 import (
    generate_sales_performance_dashboard_2026,
    validate_benchmark_workbook,
    BENCHMARK_TRANSACTIONS,
)
from backend.tools.excel_analyzer import (
    execute_excel_analysis,
    active_excel_context,
)
from backend.tools.excel_parser import parse_excel_generation_request
from backend.llm.router import (
    is_excel_generation_request,
    is_excel_formula_explanation,
    is_weather_request,
    is_standalone_math_request,
    query_router,
)


# ==============================================================================
# 1. Benchmark Workbook Generation & Validation (Phases 4 & 5)
# ==============================================================================

def test_benchmark_workbook_generation_and_sheets():
    """Verify benchmark generates a valid XLSX with exactly the 4 required sheets."""
    artifact = generate_sales_performance_dashboard_2026(unique_filename=True)

    assert artifact["status"] == "ready"
    assert artifact["sheet_count"] == 4
    expected_sheets = ["Dashboard", "Sales Data", "Analysis", "Lists"]
    assert artifact["sheet_names"] == expected_sheets

    file_path = artifact["file_path"]
    assert os.path.exists(file_path)
    assert artifact["file_size"] > 0

    wb = openpyxl.load_workbook(file_path, data_only=True)
    assert wb.sheetnames == expected_sheets
    wb.close()


def test_benchmark_sales_data_rows_and_formulas():
    """
    Verify Sales Data contains exactly 30 transactions, unique IDs,
    and active Excel formulas in columns H, J, K.
    """
    artifact = generate_sales_performance_dashboard_2026(unique_filename=True)
    file_path = artifact["file_path"]

    wb_f = openpyxl.load_workbook(file_path, data_only=False)
    ws_data = wb_f["Sales Data"]

    # Verify exactly 30 transactions (rows 2 to 31)
    txn_ids = []
    for r in range(2, 32):
        t_id = str(ws_data.cell(row=r, column=1).value)
        txn_ids.append(t_id)

        # Check formulas
        h_formula = str(ws_data.cell(row=r, column=8).value)
        j_formula = str(ws_data.cell(row=r, column=10).value)
        k_formula = str(ws_data.cell(row=r, column=11).value)

        assert h_formula == f"=F{r}*G{r}", f"Row {r} total sales formula mismatch: {h_formula}"
        assert j_formula == f"=H{r}-I{r}", f"Row {r} profit formula mismatch: {j_formula}"
        assert k_formula == f"=IFERROR(J{r}/H{r},0)", f"Row {r} margin formula mismatch: {k_formula}"

        # Confirm not escaped with apostrophe
        assert not h_formula.startswith("'=")
        assert not j_formula.startswith("'=")
        assert not k_formula.startswith("'=")

    assert len(txn_ids) == 30
    assert len(set(txn_ids)) == 30, "Transaction IDs must be unique across all 30 transactions."
    wb_f.close()


def test_benchmark_independent_calculation_reconciliation():
    """
    Phase 5 Reconciliation:
    Independently calculate totals in Python and verify reconciliation.
    """
    artifact = generate_sales_performance_dashboard_2026(unique_filename=True)
    validation = artifact["validation"]

    assert validation["status"] == "passed"
    rec = validation["reconciliation"]

    # Calculate expected numbers independently from raw transaction tuples
    expected_sales = sum(t[5] * t[6] for t in BENCHMARK_TRANSACTIONS)
    expected_profit = sum((t[5] * t[6]) - t[7] for t in BENCHMARK_TRANSACTIONS)
    expected_units = sum(t[5] for t in BENCHMARK_TRANSACTIONS)

    assert rec["total_revenue"] == expected_sales
    assert rec["total_profit"] == expected_profit
    assert rec["total_units_sold"] == expected_units
    assert rec["reconciles_regional"] is True
    assert rec["reconciles_product"] is True
    assert len(validation["formula_errors_detected"]) == 0


def test_benchmark_charts_and_data_validation():
    """Verify 5 charts and dropdown data validations exist."""
    artifact = generate_sales_performance_dashboard_2026(unique_filename=True)
    file_path = artifact["file_path"]

    wb = openpyxl.load_workbook(file_path)
    ws_analysis = wb["Analysis"]
    ws_data = wb["Sales Data"]

    # 5 charts on Analysis
    assert len(ws_analysis._charts) >= 5

    # Data validation dropdowns on Sales Data
    assert len(ws_data.data_validations.dataValidation) >= 2

    # Conditional formatting rules on Sales Data
    assert len(ws_data.conditional_formatting) >= 1

    wb.close()


# ==============================================================================
# 2. Spreadsheet Analysis Operations (Phases 2 & 6)
# ==============================================================================

@pytest.fixture
def benchmark_artifact():
    art = generate_sales_performance_dashboard_2026(unique_filename=True)
    active_excel_context.set({
        "doc_id": art["artifact_id"],
        "file_path": art["file_path"],
        "filename": art["filename"],
        "has_spreadsheet": True
    })
    return art


def test_analyzer_top_n_ranking(benchmark_artifact):
    """Verify top 5 ranking operation on sales and profit."""
    res_top_profit = execute_excel_analysis(
        operation="top_n",
        column="Profit",
        n=5,
        sheet_name="Sales Data"
    )
    assert res_top_profit["status"] == "success"
    records = res_top_profit["records"]
    assert len(records) == 5
    assert records[0]["rank"] == 1
    # Verify descending profit
    assert records[0]["Profit"] >= records[1]["Profit"]
    assert records[1]["Profit"] >= records[2]["Profit"]


def test_analyzer_monthly_summary(benchmark_artifact):
    """Verify monthly grouping across calendar months."""
    res_monthly = execute_excel_analysis(
        operation="monthly_summary",
        sheet_name="Sales Data"
    )
    assert res_monthly["status"] == "success"
    breakdown = res_monthly["monthly_breakdown"]
    assert len(breakdown) >= 4  # Jan, Feb, Mar, Apr 2026
    months = [m["month"] for m in breakdown]
    assert any("Jan" in m for m in months)
    assert any("Feb" in m for m in months)
    assert any("Mar" in m for m in months)
    assert any("Apr" in m for m in months)


def test_analyzer_missing_duplicates_audit(benchmark_artifact):
    """Verify missing values and duplicate records detection."""
    res_audit = execute_excel_analysis(
        operation="missing_duplicates",
        sheet_name="Sales Data"
    )
    assert res_audit["status"] == "success"
    assert res_audit["total_rows"] == 30
    assert res_audit["total_duplicate_rows"] == 0
    assert res_audit["duplicate_ids_count"] == 0
    assert res_audit["is_clean"] is True


def test_analyzer_natural_language_queries(benchmark_artifact):
    """Verify natural language queries return grounded answers with calculations."""
    # 1. Total revenue
    q1 = execute_excel_analysis(operation="query", query="What is the total sales revenue?")
    assert q1["status"] == "success"
    assert "₹3,380,583.00" in q1["answer"]

    # 2. Highest revenue region
    q2 = execute_excel_analysis(operation="query", query="Which region generated the highest revenue?")
    assert q2["status"] == "success"
    assert "South" in q2["answer"] or "North" in q2["answer"]

    # 3. Top 5 profitable
    q3 = execute_excel_analysis(operation="query", query="Find the top five most profitable transactions")
    assert q3["status"] == "success"
    assert "Rank" in q3["answer"] or "#1" in q3["answer"]


# ==============================================================================
# 3. AI Agent Tool Routing (Phase 3)
# ==============================================================================

def test_routing_benchmark_generation():
    """Verify prompt asking for benchmark dashboard routes to excel_generation."""
    p1 = "Generate the Sales Performance Analysis Dashboard — 2026"
    assert is_excel_generation_request(p1) is True
    parsed = parse_excel_generation_request(p1)
    assert parsed.dataset_type == "sales_dashboard_2026"
    assert parsed.data_rows == 30
    assert parsed.data_columns == 11


def test_routing_formula_explanation():
    """Verify formula explanation questions are never misclassified as generation."""
    f1 = "Explain the formula =IFERROR(J2/H2,0)"
    f2 = "How does XLOOKUP work in Excel?"
    f3 = "What does the formula =SUMIFS(H2:H31, B2:B31, '>2026-01-01') do?"

    assert is_excel_formula_explanation(f1) is True
    assert is_excel_formula_explanation(f2) is True
    assert is_excel_formula_explanation(f3) is True

    # Critical requirement: MUST NOT route to workbook generation!
    assert is_excel_generation_request(f1) is False
    assert is_excel_generation_request(f2) is False
    assert is_excel_generation_request(f3) is False


@pytest.mark.asyncio
async def test_routing_full_matrix():
    """Verify query router correctly routes the full matrix of user intents."""
    # 1. Formula explainer -> general
    r1, badge1 = await query_router.route_query("Explain this Excel formula: =INDEX(MATCH())")
    assert r1 == "general"
    assert "Formula Explainer" in badge1

    # 2. Weather -> general
    r2, badge2 = await query_router.route_query("What is the weather in Delhi today?")
    assert r2 == "general"
    assert "Live Weather" in badge2

    # 3. Calculator -> general
    r3, badge3 = await query_router.route_query("What is 4500 * 12?")
    assert r3 == "general"
    assert "Calculator" in badge3

    # 4. Excel generator -> excel_generation
    r4, badge4 = await query_router.route_query("Create an Excel workbook with 30 sales transactions")
    assert r4 == "excel_generation"

    # 5. Active spreadsheet loaded + sales question -> Spreadsheet Analysis
    active_excel_context.set({"has_spreadsheet": True, "doc_ids": ["test_sheet_123"]})
    r5, badge5 = await query_router.route_query("What is the total sales revenue?")
    assert r5 == "document"
    assert "Spreadsheet Analysis" in badge5

    # 6. PDF document question -> From document
    active_excel_context.set(None)
    r6, badge6 = await query_router.route_query("What are the key findings mentioned in the PDF report?")
    assert r6 == "document"
    assert "From document" in badge6


# ==============================================================================
# 4. Security: CWE-1236 Formula Injection Defense (Phase 8)
# ==============================================================================

def test_cwe_1236_sanitization_rules():
    """
    Verify:
    1. Safe intentional formulas are preserved when allow_formula=True.
    2. Malicious DDE execution payloads are strictly neutralized.
    3. Untrusted user strings starting with formula characters are escaped.
    """
    # Safe formula
    safe_f = "=F2*G2"
    assert sanitize_cell_value(safe_f, allow_formula=True) == safe_f

    safe_f2 = "=IFERROR(J2/H2, 0)"
    assert sanitize_cell_value(safe_f2, allow_formula=True) == safe_f2

    # Malicious DDE injection attacks MUST be escaped even with allow_formula=True
    bad_cmd = "=cmd|' /C calc'!A0"
    assert sanitize_cell_value(bad_cmd, allow_formula=True) == f"'{bad_cmd}"

    bad_ps = "=powershell -e abc"
    assert sanitize_cell_value(bad_ps, allow_formula=True) == f"'{bad_ps}"

    # Untrusted user text (allow_formula=False default)
    user_input = "=SUM(1,2)"
    assert sanitize_cell_value(user_input, allow_formula=False) == f"'{user_input}"

    # Plain text and numbers
    assert sanitize_cell_value(12345) == 12345
    assert sanitize_cell_value(8500.50) == 8500.50
    assert sanitize_cell_value("Normal Text") == "Normal Text"
